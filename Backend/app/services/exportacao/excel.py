"""Geração de relatórios em Excel (openpyxl via pandas).

A primeira aba ("Resumo") traz o cabeçalho do relatório — hospital, quem gerou
e quando — e os indicadores. Depois vêm a tabela principal completa e as abas
de detalhe, com cabeçalho fixo, filtro e formato numérico nas colunas.
"""

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.exportacao.dados import CASAS_DECIMAIS, COLUNAS_MOEDA, Relatorio
from app.services.exportacao.pdf import ContextoExportacao

_NAVY = "0A2540"
_CINZA = "5B6B7E"


def _nome_aba(nome: str) -> str:
    # Nomes de aba do Excel: no máximo 31 caracteres e sem alguns símbolos.
    for proibido in "[]:*?/\\":
        nome = nome.replace(proibido, "-")
    return nome[:31]


def _formatar_tabela(ws: Worksheet, df: pd.DataFrame) -> None:
    cabecalho = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor=_NAVY)
    for indice, coluna in enumerate(df.columns, start=1):
        celula = ws.cell(row=1, column=indice)
        celula.font = cabecalho
        celula.fill = fundo
        celula.alignment = Alignment(vertical="center")

        letra = get_column_letter(indice)
        maior = max([len(str(coluna))] + [len(str(v)) for v in df[coluna].head(500)])
        ws.column_dimensions[letra].width = min(max(maior + 2, 9), 60)

        if coluna in CASAS_DECIMAIS:
            casas = CASAS_DECIMAIS[coluna]
            formato = "#,##0" + ("." + "0" * casas if casas else "")
            if coluna in COLUNAS_MOEDA:
                formato = '"R$" ' + formato
            for (celula_valor,) in ws.iter_rows(min_row=2, min_col=indice, max_col=indice):
                celula_valor.number_format = formato
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions


def _aba_resumo(ws: Worksheet, relatorio: Relatorio, contexto: ContextoExportacao) -> None:
    ws["A1"] = relatorio.titulo
    ws["A1"].font = Font(bold=True, size=15, color=_NAVY)
    if relatorio.referencia:
        ws["A2"] = relatorio.referencia
        ws["A2"].font = Font(italic=True, color=_CINZA)

    linhas = [
        ("Hospital", contexto.hospital),
        ("Gerado por", contexto.usuario),
        ("Data de geração", contexto.gerado_em.strftime("%d/%m/%Y %H:%M")),
        ("", ""),
        *relatorio.indicadores,
    ]
    for deslocamento, (rotulo, valor) in enumerate(linhas):
        linha = 4 + deslocamento
        ws.cell(row=linha, column=1, value=rotulo).font = Font(color=_CINZA)
        ws.cell(row=linha, column=2, value=valor).font = Font(bold=True, color=_NAVY)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 48


def gerar_xlsx(relatorio: Relatorio, contexto: ContextoExportacao) -> bytes:
    buffer = BytesIO()
    abas = {relatorio.nome_tabela: relatorio.tabela, **relatorio.secoes, **relatorio.abas_extras}
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        resumo = writer.book.create_sheet("Resumo", 0)
        _aba_resumo(resumo, relatorio, contexto)
        for nome, df in abas.items():
            aba = _nome_aba(nome)
            df.to_excel(writer, sheet_name=aba, index=False)
            _formatar_tabela(writer.sheets[aba], df)
    return buffer.getvalue()
