"""Testes da camada de dados e dos geradores de relatório (RF10, RN04)."""

from datetime import date, datetime
from io import BytesIO

import base64
import re
import zlib

import openpyxl
import pandas as pd

from app.models.classificacao_abc import ClassificacaoABC
from app.models.consumo_tratado import ConsumoTratado
from app.models.itens import Item
from app.models.modelo_treinado import ModeloTreinado
from app.models.previsoes import Previsao
from app.services.exportacao import dados
from app.services.exportacao.excel import gerar_xlsx
from app.services.exportacao.pdf import LIMITE_LINHAS_PDF, ContextoExportacao, gerar_pdf

CONTEXTO = ContextoExportacao(
    hospital="Hospital & Maternidade Teste", usuario="Ana Souza", gerado_em=datetime(2026, 10, 20, 14, 30)
)


def _item(db, empresa, codigo="MED001", **attrs) -> Item:
    item = Item(empresa_id=empresa.id, codigo_item=codigo, descricao_item=f"Item {codigo}", **attrs)
    db.add(item)
    db.flush()
    return item


def _tratado(db, empresa, item, periodo, qtd, valor, setor="Central") -> None:
    db.add(ConsumoTratado(
        empresa_id=empresa.id, item_id=item.id, periodo=periodo,
        quantidade_total=qtd, valor_total=valor, setor=setor,
    ))


def _previsoes(db, empresa, item, valores=(10.0, 20.0, 30.0), *, algoritmo="holt_winters", mae=4.0):
    """Modelo ativo da empresa (criado na 1ª chamada) com a previsão mensal do item."""
    modelo = (
        db.query(ModeloTreinado)
        .filter(ModeloTreinado.empresa_id == empresa.id, ModeloTreinado.ativo.is_(True))
        .first()
    )
    if modelo is None:
        modelo = ModeloTreinado(
            empresa_id=empresa.id, algoritmo="melhor_por_item", ativo=True,
            treinado_em=datetime(2026, 10, 1, 8, 0), parametros={"itens": {}},
        )
        db.add(modelo)
        db.flush()
    modelo.parametros = {
        "itens": {**modelo.parametros["itens"], str(item.id): {"algoritmo": algoritmo, "mae": mae}}
    }
    for i, v in enumerate(valores):
        db.add(Previsao(
            empresa_id=empresa.id, modelo_id=modelo.id, item_id=item.id,
            periodo=date(2024, 7 + i, 1), quantidade_prevista=v,
            intervalo_inferior=v * 0.8, intervalo_superior=v * 1.2,
        ))
    # Referência: último mês com dados = junho/2024.
    _tratado(db, empresa, item, date(2024, 6, 1), 1, 1)


def _abc(db, empresa, item, classe) -> None:
    db.add(ClassificacaoABC(
        empresa_id=empresa.id, item_id=item.id, classe=classe, valor_acumulado_percentual=0.5
    ))


# --------------------------------------------------------------------------- #
# Camada de dados
# --------------------------------------------------------------------------- #
def test_relatorio_previsoes_uma_linha_por_material_com_os_3_meses(db_session, empresa):
    item = _item(db_session, empresa, unidade_medida="AMP")
    _previsoes(db_session, empresa, item, (10.0, 20.0, 30.0))
    _abc(db_session, empresa, item, "A")
    db_session.commit()

    rel = dados.montar_relatorio_previsoes(db_session, empresa.id)
    assert list(rel.tabela.columns) == dados.COLUNAS_PREVISOES
    assert len(rel.tabela) == 1
    linha = rel.tabela.iloc[0]
    assert (linha["Unid."], linha["Curva ABC"]) == ("AMP", "A")
    assert (linha["Mês 1"], linha["Mês 2"], linha["Mês 3"], linha["Total 3 meses"]) == (10.0, 20.0, 30.0, 60.0)
    assert ("Curva A", "1") in rel.indicadores
    assert ("Modelo treinado em", "01/10/2026") in rel.indicadores
    assert "Mês 1 = 07/2024" in rel.referencia and "06/2024" in rel.referencia


def test_relatorio_previsoes_detalhe_mensal_no_excel(db_session, empresa):
    item = _item(db_session, empresa)
    _previsoes(db_session, empresa, item, (100.0, 110.0, 120.0), algoritmo="holt_winters", mae=4.0)
    db_session.commit()

    detalhe = dados.montar_relatorio_previsoes(db_session, empresa.id).abas_extras["Detalhe mensal"]
    assert list(detalhe.columns) == dados.COLUNAS_PREVISAO_MENSAL
    assert list(detalhe["Competência"]) == ["07/2024", "08/2024", "09/2024"]
    assert set(detalhe["Modelo"]) == {"Holt-Winters"}
    primeira = detalhe.iloc[0]
    assert (primeira["Mínimo"], primeira["Máximo"], primeira["Erro médio (MAE)"]) == (80.0, 120.0, 4.0)


def test_relatorio_previsoes_ordena_por_curva_abc(db_session, empresa):
    c = _item(db_session, empresa, "C1")
    a = _item(db_session, empresa, "A1")
    for item, classe in ((c, "C"), (a, "A")):
        _previsoes(db_session, empresa, item)
        _abc(db_session, empresa, item, classe)
    db_session.commit()
    rel = dados.montar_relatorio_previsoes(db_session, empresa.id)
    assert list(rel.tabela["Código"]) == ["A1", "C1"]


def test_relatorio_previsoes_isolado_por_empresa(db_session, empresa, empresa_secundaria):
    _previsoes(db_session, empresa, _item(db_session, empresa))
    _previsoes(db_session, empresa_secundaria, _item(db_session, empresa_secundaria, "OUTRO"))
    db_session.commit()
    assert len(dados.montar_relatorio_previsoes(db_session, empresa.id).tabela) == 1


def test_relatorio_previsoes_vazio(db_session, empresa):
    rel = dados.montar_relatorio_previsoes(db_session, empresa.id)
    assert rel.tabela.empty
    assert "Ainda não há" in rel.referencia


def test_relatorio_consumo_resume_por_material_e_setor(db_session, empresa):
    caro = _item(db_session, empresa, "CARO", grupo_material="Medicamentos")
    barato = _item(db_session, empresa, "BARATO")
    _tratado(db_session, empresa, caro, date(2024, 1, 1), 10, 600, setor="UTI")
    _tratado(db_session, empresa, caro, date(2024, 2, 1), 10, 200, setor="Farmácia")
    _tratado(db_session, empresa, barato, date(2024, 2, 1), 5, 200, setor="UTI")
    db_session.commit()

    rel = dados.montar_relatorio_consumo(db_session, empresa.id)
    assert list(rel.tabela.columns) == dados.COLUNAS_CONSUMO
    assert list(rel.tabela["Código"]) == ["CARO", "BARATO"]  # do maior para o menor valor
    assert list(rel.tabela["% do valor"]) == [80.0, 20.0]
    assert rel.tabela.iloc[0]["Grupo"] == "Medicamentos"

    setores = rel.secoes["Consumo por setor"]
    assert list(setores["Setor"]) == ["UTI", "Farmácia"]
    assert len(rel.abas_extras["Detalhe mensal"]) == 3  # item x mês, somando os setores
    assert ("Valor consumido", "R$ 1.000,00") in rel.indicadores
    assert rel.referencia == "Competências de 01/2024 a 02/2024"


def test_relatorio_consumo_filtra_data(db_session, empresa):
    item = _item(db_session, empresa)
    _tratado(db_session, empresa, item, date(2024, 1, 1), 10, 100)
    _tratado(db_session, empresa, item, date(2024, 6, 1), 20, 200)
    db_session.commit()
    rel = dados.montar_relatorio_consumo(db_session, empresa.id, data_inicio=date(2024, 3, 1))
    assert rel.tabela.iloc[0]["Valor (R$)"] == 200.0


def test_fmt_numero_pt_br():
    assert dados.fmt_numero(1234567.891, 2) == "1.234.567,89"
    assert dados.fmt_moeda(12.5) == "R$ 12,50"


# --------------------------------------------------------------------------- #
# Geradores
# --------------------------------------------------------------------------- #
def _relatorio_grande(n: int) -> dados.Relatorio:
    tabela = pd.DataFrame(
        [[f"M{i}", f"Material <{i}> & cia", "UN", "—", "A", 1.0, 10.0 * i, 1.0] for i in range(n)],
        columns=dados.COLUNAS_CONSUMO,
    )
    return dados.Relatorio(titulo="Relatório de Consumo", tabela=tabela, indicadores=[("Materiais", str(n))])


def _texto_pdf(conteudo: bytes) -> tuple[str, int]:
    """Texto bruto (streams descomprimidos) e nº de páginas, sem depender de leitor de PDF."""
    partes = []
    for stream in re.findall(rb"stream\r?\n(.*?)endstream", conteudo, re.S):
        bruto = stream.strip()
        try:  # o reportlab usa /ASCII85Decode + /FlateDecode
            bruto = base64.a85decode(bruto, adobe=True) if bruto.endswith(b"~>") else bruto
            partes.append(zlib.decompress(bruto).decode("latin-1"))
        except (ValueError, zlib.error):
            partes.append(stream.decode("latin-1"))
    paginas = len(re.findall(rb"/Type /Page\b", conteudo))
    return "\n".join(partes), paginas


def test_gerar_pdf_tem_cabecalho_e_limita_linhas():
    conteudo = gerar_pdf(_relatorio_grande(LIMITE_LINHAS_PDF + 50), CONTEXTO)
    assert conteudo[:4] == b"%PDF"
    texto, paginas = _texto_pdf(conteudo)
    assert paginas > 1
    assert "Maternidade Teste" in texto
    assert "Ana Souza" in texto
    assert "20/10/2026" in texto
    assert f"gina 1 de {paginas}" in texto
    assert "Excel" in texto
    assert f"(M{LIMITE_LINHAS_PDF - 1})" in texto
    assert f"(M{LIMITE_LINHAS_PDF + 10})" not in texto


def test_gerar_pdf_vazio_nao_quebra(db_session, empresa):
    rel = dados.montar_relatorio_previsoes(db_session, empresa.id)
    assert gerar_pdf(rel, CONTEXTO)[:4] == b"%PDF"


def test_gerar_xlsx_tem_resumo_e_lista_completa():
    conteudo = gerar_xlsx(_relatorio_grande(LIMITE_LINHAS_PDF + 50), CONTEXTO)
    planilha = openpyxl.load_workbook(BytesIO(conteudo))
    assert planilha.sheetnames[0] == "Resumo"
    valores = [c.value for linha in planilha["Resumo"].iter_rows() for c in linha]
    assert "Hospital & Maternidade Teste" in valores and "Ana Souza" in valores
    materiais = planilha["Materiais"]
    assert materiais.max_row == LIMITE_LINHAS_PDF + 51  # cabeçalho + todas as linhas
    assert materiais["G2"].number_format.startswith('"R$"')
